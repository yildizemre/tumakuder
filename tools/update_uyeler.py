# -*- coding: utf-8 -*-
"""Update assets/data/uyeler.json from Web için Üye Firma Bilgileri.xlsx"""
import json
import re
from pathlib import Path

import openpyxl

root = Path(__file__).resolve().parents[1]
xlsx = root / "Web için Üye Firma Bilgileri.xlsx"
uyeler_path = root / "assets" / "data" / "uyeler.json"


def norm(s):
    if s is None:
        return ""
    s = str(s).replace("\xa0", " ")
    s = re.sub(r"\s+", " ", s).strip().upper()
    tr = str.maketrans("İIĞÜŞÖÇ", "IIGUSOC")
    return s.translate(tr).replace("İ", "I")


def clean(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s).replace("\xa0", " ")).strip()


def merge_contact(entry, prev):
    if not prev:
        return
    for k in ("tel", "eposta", "site"):
        if prev.get(k):
            entry[k] = prev[k]


old = json.loads(uyeler_path.read_text(encoding="utf-8"))
old_asil_by_kod = {str(u["kod"]).strip(): u for u in old.get("asil", []) if u.get("kod")}
old_asil_by_name = {norm(u.get("firma")): u for u in old.get("asil", [])}
old_fahri_by_name = {norm(u.get("firma")): u for u in old.get("fahri", [])}

wb = openpyxl.load_workbook(xlsx, data_only=True)

asil = []
for row in list(wb["Asil Üye Firmalar"].iter_rows(values_only=True))[1:]:
    if not row or not row[1]:
        continue
    kod = str(int(row[0])) if isinstance(row[0], (int, float)) else str(row[0]).strip()
    firma = clean(row[1])
    sehir = clean(row[2])
    entry = {"kod": kod, "firma": firma, "sehir": sehir}
    merge_contact(entry, old_asil_by_kod.get(kod) or old_asil_by_name.get(norm(firma)))
    asil.append(entry)

fahri = []
for row in list(wb["Fahri Üye Firmalar"].iter_rows(values_only=True))[1:]:
    if not row or not row[0]:
        continue
    firma = clean(row[0])
    sehir = clean(row[1])
    entry = {"firma": firma, "sehir": sehir}
    merge_contact(entry, old_fahri_by_name.get(norm(firma)))
    fahri.append(entry)

old_asil_names = {norm(u["firma"]) for u in old["asil"]}
new_asil_names = {norm(u["firma"]) for u in asil}
old_fahri_names = {norm(u["firma"]) for u in old["fahri"]}
new_fahri_names = {norm(u["firma"]) for u in fahri}

print(f"Asil: {len(old['asil'])} -> {len(asil)}")
print(f"  removed: {len(old_asil_names - new_asil_names)}")
print(f"  added:   {len(new_asil_names - old_asil_names)}")
print(f"Fahri: {len(old['fahri'])} -> {len(fahri)}")
print(f"  removed: {len(old_fahri_names - new_fahri_names)}")
print(f"  added:   {len(new_fahri_names - old_fahri_names)}")

print("--- asil added ---")
for n in sorted(new_asil_names - old_asil_names):
    for u in asil:
        if norm(u["firma"]) == n:
            print(" +", u["kod"], u["firma"][:70], u["sehir"])
            break

print("--- asil removed ---")
for n in sorted(old_asil_names - new_asil_names):
    for u in old["asil"]:
        if norm(u["firma"]) == n:
            print(" -", u.get("kod"), u["firma"][:70], u["sehir"])
            break

print("--- fahri added ---")
for n in sorted(new_fahri_names - old_fahri_names):
    for u in fahri:
        if norm(u["firma"]) == n:
            print(" +", u["firma"][:70], u["sehir"])
            break

print("--- fahri removed ---")
for n in sorted(old_fahri_names - new_fahri_names):
    for u in old["fahri"]:
        if norm(u["firma"]) == n:
            print(" -", u["firma"][:70], u["sehir"])
            break

asil_with = sum(1 for u in asil if u.get("tel") or u.get("eposta") or u.get("site"))
fahri_with = sum(1 for u in fahri if u.get("tel") or u.get("eposta") or u.get("site"))
print(f"Contact preserved: asil {asil_with}/{len(asil)}, fahri {fahri_with}/{len(fahri)}")

uyeler_path.write_text(
    json.dumps({"asil": asil, "fahri": fahri}, ensure_ascii=False, indent=1) + "\n",
    encoding="utf-8",
)
print("Wrote", uyeler_path)
